"""Export calculation memorial HTML to DOCX with editable Word equations."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import ClassVar

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_EQUATION_MARKER = re.compile(r"OSA_EQUATION_OBJECT_\d{4}_END")
_OMML_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/math"


@dataclass
class _Node:
    tag: str
    children: list[_Node | str] = field(default_factory=list)


class _MemorialParser(HTMLParser):
    """Build a small tree for the limited HTML emitted by the memorial."""

    _VOID_TAGS: ClassVar[set[str]] = {"br", "hr", "img", "meta", "link", "input"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = _Node("root")
        self._stack = [self.root]

    def handle_starttag(self, tag: str, attrs) -> None:
        node = _Node(tag.lower())
        self._stack[-1].children.append(node)
        if tag.lower() not in self._VOID_TAGS:
            self._stack.append(node)

    def handle_startendtag(self, tag: str, attrs) -> None:
        self.handle_starttag(tag, attrs)
        if tag.lower() not in self._VOID_TAGS:
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        for index in range(len(self._stack) - 1, 0, -1):
            if self._stack[index].tag == tag:
                del self._stack[index:]
                break

    def handle_data(self, data: str) -> None:
        if data:
            self._stack[-1].children.append(data)


def _omml_element(latex: str):
    """Convert one supported memorial LaTeX expression into an OMML element."""
    from latex2mathml.converter import convert as latex_to_mathml
    from mathml2omml import convert as mathml_to_omml

    try:
        omml = mathml_to_omml(latex_to_mathml(latex))
        # mathml2omml emits the m: prefix without declaring its namespace.
        xml = omml.replace(
            "<m:oMath>",
            f'<m:oMath xmlns:m="{_OMML_NAMESPACE}">',
            1,
        )
        equation = parse_xml(xml)
        for radical in equation.iter(f"{{{_OMML_NAMESPACE}}}rad"):
            # A raiz quadrada precisa declarar que o índice está oculto.
            # Sem radPr/deg, alguns leitores (incluindo o Writer) exibem
            # caixas e fragmentos em vez do radicando.
            radical_properties = OxmlElement("m:radPr")
            degree_hidden = OxmlElement("m:degHide")
            degree_hidden.set(qn("m:val"), "1")
            radical_properties.append(degree_hidden)
            degree = OxmlElement("m:deg")
            radical.insert(0, radical_properties)
            radical.insert(1, degree)
        return equation
    except Exception as exc:
        raise ValueError(f"Não foi possível converter a equação LaTeX {latex!r} para OMML: {exc}") from exc


def _append_inline(paragraph, children, equations: dict[str, str], bold: bool = False) -> None:
    for child in children:
        if isinstance(child, str):
            cursor = 0
            for match in _EQUATION_MARKER.finditer(child):
                text = child[cursor:match.start()]
                if text:
                    run = paragraph.add_run(text)
                    run.bold = bold
                latex = equations.get(match.group())
                if latex is None:
                    run = paragraph.add_run(match.group())
                    run.bold = bold
                else:
                    paragraph._p.append(_omml_element(latex))
                cursor = match.end()
            tail = child[cursor:]
            if tail:
                run = paragraph.add_run(tail)
                run.bold = bold
        elif child.tag == "br":
            paragraph.add_run().add_break()
        elif child.tag in {"b", "strong"}:
            _append_inline(paragraph, child.children, equations, bold=True)
        else:
            _append_inline(paragraph, child.children, equations, bold=bold)


def _table_rows(node: _Node) -> list[_Node]:
    rows = []
    for child in node.children:
        if isinstance(child, _Node):
            if child.tag == "tr":
                rows.append(child)
            elif child.tag in {"thead", "tbody", "tfoot"}:
                rows.extend(_table_rows(child))
    return rows


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", "80"), ("start", "0"), ("bottom", "80"), ("end", "0")):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")


def _remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def _export_node(document, node: _Node, equations: dict[str, str]) -> None:
    if node.tag in {"h2", "h3"}:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(11 if node.tag == "h2" else 16)
        paragraph.paragraph_format.space_after = Pt(8 if node.tag == "h2" else 6)
        _append_inline(paragraph, node.children, equations, bold=True)
        for run in paragraph.runs:
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(13.5 if node.tag == "h2" else 11.25)
            run.font.color.rgb = RGBColor.from_string("24292F" if node.tag == "h2" else "57606A")
    elif node.tag == "p":
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _append_inline(paragraph, node.children, equations)
    elif node.tag == "table":
        rows = _table_rows(node)
        table_rows = []
        for row in rows:
            cells = [
                child for child in row.children
                if isinstance(child, _Node) and child.tag in {"td", "th"}
            ]
            if cells:
                table_rows.append(cells)
        if not table_rows:
            return
        table = document.add_table(rows=len(table_rows), cols=2)
        table.autofit = False
        table.columns[0].width = Inches(3.8)
        table.columns[1].width = Inches(2.7)
        _remove_table_borders(table)
        for row_index, cells in enumerate(table_rows):
            for column_index in range(2):
                cell = table.cell(row_index, column_index)
                cell.width = Inches(3.8 if column_index == 0 else 2.7)
                _set_cell_margins(cell)
                if column_index >= len(cells):
                    continue
                _append_inline(
                    cell.paragraphs[0],
                    cells[column_index].children,
                    equations,
                    bold=column_index == 0 or cells[column_index].tag == "th",
                )
    elif node.tag == "div":
        for child in node.children:
            if isinstance(child, _Node):
                _export_node(document, child, equations)


def export_calculation_memorial_docx(
    html: str,
    equation_latex: dict[str, str],
    destination: str | Path,
) -> Path:
    """Write memorial text, tables, and native editable equations to a DOCX."""
    parser = _MemorialParser()
    parser.feed(html)
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(9.75)
    normal_style.font.color.rgb = RGBColor.from_string("57606A")
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    document.core_properties.title = "Memorial de cálculo de elemento estrutural"
    document.core_properties.subject = "Verificações estruturais conforme ABNT NBR 8800"
    for child in parser.root.children:
        if isinstance(child, _Node):
            _export_node(document, child, equation_latex)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination
