import os
from zipfile import ZipFile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QFileDialog, QLabel, QMessageBox, QPushButton

from osa.ui.calculation_memorial_dialog import CalculationMemorialDialog, MemorialExportSuccessDialog
from osa.ui.memorial_docx_export import export_calculation_memorial_docx


def test_docx_export_preserves_content_and_inserts_editable_omml(tmp_path):
    marker = "OSA_EQUATION_OBJECT_0000_END"
    html = (
        "<div><h2>Memorial de cálculo</h2><h3>Dados da barra</h3>"
        "<table><tr><td>Perfil</td><td>W 200</td></tr></table>"
        f"<p>Força resistente: {marker}</p></div>"
    )
    destination = export_calculation_memorial_docx(
        html,
        {marker: r"\sqrt{\dfrac{A_g \cdot f_y}{\gamma_{a1}}}"},
        tmp_path / "memorial.docx",
    )

    document = Document(destination)
    assert document.paragraphs[0].text == "Memorial de cálculo"
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "Perfil"
    assert document.tables[0].cell(0, 1).text == "W 200"
    assert document.paragraphs[-1].text.startswith("Força resistente:")
    assert document.paragraphs[-1].alignment == WD_ALIGN_PARAGRAPH.LEFT
    with ZipFile(destination) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        assert "<m:oMath" in document_xml
        assert "<m:f>" in document_xml
        assert "<m:rad" in document_xml
        assert "<m:degHide m:val=\"1\"" in document_xml
        assert 'w:tblBorders' in document_xml and 'w:val="nil"' in document_xml
        assert archive.testzip() is None


def test_export_button_saves_docx_to_user_selected_path(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    destination = tmp_path / "memorial-gerado.docx"
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        lambda *_args, **_kwargs: (str(destination), "Documento Word (*.docx)"),
    )
    monkeypatch.setattr(QMessageBox, "warning", lambda *_args, **_kwargs: QMessageBox.StandardButton.Ok)
    success_messages = []
    monkeypatch.setattr(
        MemorialExportSuccessDialog,
        "exec",
        lambda self: success_messages.append(self.findChild(QLabel).text()),
    )
    dialog = CalculationMemorialDialog({})
    button = dialog.findChild(QPushButton, "exportMemorial")
    assert button is not None
    button.click()

    assert destination.is_file()
    assert len(Document(destination).paragraphs) > 0
    assert success_messages == ["Memorial exportado com sucesso"]
    dialog.close()
    app.processEvents()


def test_memorial_dialog_keeps_original_width():
    app = QApplication.instance() or QApplication([])
    dialog = CalculationMemorialDialog({})

    assert dialog.width() == 700

    dialog.close()
    app.processEvents()


def test_memorial_equation_paragraphs_are_centered():
    equations = {}
    memory = {
        "profile_type": "W Laminado",
        "axial_tension": {
            "verification": {
                "status": "checked",
                "parameters": {
                    "Ag_cm2": 10.0,
                    "An_cm2": 9.0,
                    "Ae_cm2": 9.0,
                    "Ct": 1.0,
                    "fy_kn_cm2": 34.5,
                    "fu_kn_cm2": 45.0,
                    "gamma_a1": 1.1,
                    "gamma_a2": 1.35,
                    "Nt_Rd1_kn": 313.64,
                    "Nt_Rd2_kn": 300.0,
                    "Nt_Rd_kn": 300.0,
                },
                "applied_kn": 100.0,
                "utilization": 1.0 / 3.0,
                "critical_combination": "Combinação 01",
            }
        },
    }

    html = CalculationMemorialDialog._render(memory, {}, equation_latex=equations)

    assert equations
    assert 'text-align:center' in html

    app = QApplication.instance() or QApplication([])
    dialog = CalculationMemorialDialog(memory)
    block = dialog.content.document().begin()
    equation_blocks = []
    while block.isValid():
        if "\ufffc" in block.text():
            equation_blocks.append(block)
        block = block.next()
    assert equation_blocks
    assert all(
        block.blockFormat().alignment() == Qt.AlignmentFlag.AlignHCenter
        and block.blockFormat().leftMargin() == 0.0
        and block.blockFormat().textIndent() == 0.0
        for block in equation_blocks
    )
    dialog.close()
    app.processEvents()
